from pathlib import Path
from typing import Literal, Optional
import sqlite3
import json
import csv
import io
import re
from datetime import datetime, timezone
import calendar
from uuid import uuid4

from fastapi import APIRouter, File, Form, HTTPException, UploadFile
from pydantic import BaseModel, Field


router = APIRouter(prefix="/api/projects", tags=["Projects"])

BASE_DIR = Path(__file__).resolve().parents[3]
DB_PATH = BASE_DIR / "projects.db"
UPLOAD_DIR = BASE_DIR / "uploads" / "evidence"


def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    return conn


# ============================================================
# UNIVERSAL PROJECT MEASUREMENT
# ============================================================

PROJECT_TYPES = {
    "Infrastructure": ["roads", "transport", "construction", "urban infrastructure"],
    "Healthcare": ["health", "healthcare", "medical"],
    "Education": ["education", "school", "higher education"],
    "Welfare": ["welfare", "social", "benefit"],
    "Digital": ["digital", "it", "e-governance", "technology"],
    "Environment": ["environment", "ecology", "forestry", "climate"],
    "Agriculture": ["agriculture", "farming", "irrigation"],
    "Housing": ["housing", "homes", "residential"],
    "Water": ["water", "drinking water", "sanitation"],
    "Energy": ["energy", "solar", "power", "electricity"],
    "Transport": ["transport", "metro", "rail", "transit"],
    "Administrative": ["administrative", "governance", "public service"],
    "Other": [],
}


METRIC_TEMPLATES = {
    "Infrastructure": [
        ("Physical delivery", "Quantity completed", "quantity", 60, "units"),
        ("Milestones", "Milestones completed", "milestone", 40, "milestones"),
    ],
    "Healthcare": [
        ("Service delivery", "Facilities/services delivered", "count", 50, "units"),
        ("Coverage", "People reached", "beneficiaries", 50, "people"),
    ],
    "Education": [
        ("Delivery", "Schools/services delivered", "count", 50, "units"),
        ("Reach", "Students covered", "beneficiaries", 50, "students"),
    ],
    "Welfare": [
        ("Reach", "Beneficiaries reached", "beneficiaries", 70, "people"),
        ("Delivery", "Programme milestones completed", "milestone", 30, "milestones"),
    ],
    "Digital": [
        ("Deployment", "Modules/services deployed", "count", 50, "modules"),
        ("Adoption", "Users onboarded", "beneficiaries", 50, "users"),
    ],
    "Environment": [
        ("Restoration", "Area restored", "area", 60, "hectares"),
        ("Outcome", "Environmental targets completed", "percentage", 40, "%"),
    ],
    "Agriculture": [
        ("Coverage", "Farmers covered", "beneficiaries", 50, "farmers"),
        ("Land coverage", "Area covered", "area", 50, "hectares"),
    ],
    "Housing": [
        ("Delivery", "Houses completed", "count", 70, "houses"),
        ("Occupancy", "Houses occupied", "count", 30, "houses"),
    ],
    "Water": [
        ("Coverage", "Households connected", "count", 60, "households"),
        ("Reach", "People covered", "beneficiaries", 40, "people"),
    ],
    "Energy": [
        ("Capacity", "Capacity installed", "capacity", 60, "MW"),
        ("Reach", "Connections delivered", "count", 40, "connections"),
    ],
    "Transport": [
        ("Delivery", "Transport assets/km delivered", "quantity", 60, "units"),
        ("Milestones", "Milestones completed", "milestone", 40, "milestones"),
    ],
    "Administrative": [
        ("Service delivery", "Services/processes delivered", "count", 60, "services"),
        ("Adoption", "Users/cases served", "beneficiaries", 40, "users"),
    ],
    "Other": [
        ("Primary objective", "Primary target completed", "percentage", 100, "%"),
    ],
}


def normalize_project_type(value: Optional[str], sector: str = "") -> str:
    candidate = (value or "").strip()
    if candidate in PROJECT_TYPES:
        return candidate

    sector_text = sector.lower()
    for project_type, keywords in PROJECT_TYPES.items():
        if any(keyword in sector_text for keyword in keywords):
            return project_type
    return "Other"


def now_iso() -> str:
    return datetime.utcnow().replace(microsecond=0).isoformat() + "Z"


# ============================================================
# RISK ENGINE
# ============================================================

def calculate_risk(
    approved_cost: float,
    current_expenditure: float,
    physical_progress: float,
    expected_progress: float,
) -> dict:
    approved_cost = max(0.0, float(approved_cost))
    current_expenditure = max(0.0, float(current_expenditure))
    physical_progress = max(0.0, min(100.0, float(physical_progress)))
    expected_progress = max(0.0, min(100.0, float(expected_progress)))

    progress_gap = max(0.0, expected_progress - physical_progress)
    expenditure_percent = (
        (current_expenditure / approved_cost) * 100
        if approved_cost > 0 else 0.0
    )
    cost_gap = max(0.0, expenditure_percent - physical_progress)

    delay_score = min(100.0, progress_gap * 4.0)
    cost_score = min(100.0, cost_gap * 4.0)
    overall_score = min(100.0, (delay_score * 0.55) + (cost_score * 0.45))

    def level(score: float) -> str:
        if score >= 80:
            return "Critical"
        if score >= 65:
            return "High"
        if score >= 45:
            return "Medium"
        return "Low"

    return {
        "risk_score": round(overall_score, 2),
        "delay_risk": level(delay_score),
        "cost_risk": level(cost_score),
        "delay_score": round(delay_score, 2),
        "cost_score": round(cost_score, 2),
        "progress_gap": round(progress_gap, 2),
        "expenditure_percent": round(expenditure_percent, 2),
        "cost_gap": round(cost_gap, 2),
    }


# ============================================================
# DATABASE + SAFE MIGRATION
# ============================================================

def add_column_if_missing(conn, table: str, column: str, definition: str):
    columns = {
        row["name"] for row in conn.execute(f"PRAGMA table_info({table})").fetchall()
    }
    if column not in columns:
        conn.execute(f"ALTER TABLE {table} ADD COLUMN {column} {definition}")


def init_db():
    conn = get_db()

    conn.execute("""
        CREATE TABLE IF NOT EXISTS projects (
            id TEXT PRIMARY KEY,
            name TEXT NOT NULL,
            sector TEXT NOT NULL,
            state TEXT NOT NULL,
            location TEXT NOT NULL,
            approved_cost REAL NOT NULL DEFAULT 0,
            current_expenditure REAL NOT NULL DEFAULT 0,
            physical_progress REAL NOT NULL DEFAULT 0,
            expected_progress REAL NOT NULL DEFAULT 0,
            risk_score REAL NOT NULL DEFAULT 0,
            delay_risk TEXT NOT NULL DEFAULT 'Low',
            cost_risk TEXT NOT NULL DEFAULT 'Low',
            status TEXT NOT NULL DEFAULT 'In Progress'
        )
    """)

    # Non-destructive migration for the user's existing projects.db.
    add_column_if_missing(conn, "projects", "project_type", "TEXT NOT NULL DEFAULT 'Infrastructure'")
    add_column_if_missing(conn, "projects", "description", "TEXT NOT NULL DEFAULT ''")
    add_column_if_missing(conn, "projects", "start_date", "TEXT")
    add_column_if_missing(conn, "projects", "target_date", "TEXT")
    add_column_if_missing(conn, "projects", "progress_model", "TEXT NOT NULL DEFAULT 'universal'")

    # Migrate universal-project tables too. CREATE TABLE IF NOT EXISTS does not
    # add columns when an older version of the table already exists.
    conn.execute("""
        CREATE TABLE IF NOT EXISTS project_objectives (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            project_id TEXT NOT NULL,
            title TEXT NOT NULL,
            description TEXT NOT NULL DEFAULT '',
            weight REAL NOT NULL DEFAULT 100,
            sort_order INTEGER NOT NULL DEFAULT 0,
            FOREIGN KEY(project_id) REFERENCES projects(id) ON DELETE CASCADE
        )
    """)
    # Older experimental versions may have created these tables without the
    # columns used by the current API. Add every required column safely.
    for column, definition in [
        ("project_id", "TEXT"),
        ("title", "TEXT NOT NULL DEFAULT ''"),
        ("description", "TEXT NOT NULL DEFAULT ''"),
        ("weight", "REAL NOT NULL DEFAULT 100"),
        ("sort_order", "INTEGER NOT NULL DEFAULT 0"),
    ]:
        add_column_if_missing(conn, "project_objectives", column, definition)

    objective_columns_now = {
        row["name"] for row in conn.execute("PRAGMA table_info(project_objectives)").fetchall()
    }
    if "name" in objective_columns_now and "title" in objective_columns_now:
        conn.execute("""
            UPDATE project_objectives
            SET title = name
            WHERE name IS NOT NULL AND TRIM(name) <> ''
        """)

    conn.execute("""
        CREATE TABLE IF NOT EXISTS project_metrics (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            project_id TEXT NOT NULL,
            objective_id INTEGER,
            name TEXT NOT NULL,
            kind TEXT NOT NULL DEFAULT 'quantity',
            unit TEXT NOT NULL DEFAULT '',
            baseline REAL NOT NULL DEFAULT 0,
            target REAL NOT NULL DEFAULT 100,
            current_value REAL NOT NULL DEFAULT 0,
            expected_value REAL,
            weight REAL NOT NULL DEFAULT 100,
            evidence_required INTEGER NOT NULL DEFAULT 0,
            last_updated TEXT,
            FOREIGN KEY(project_id) REFERENCES projects(id) ON DELETE CASCADE,
            FOREIGN KEY(objective_id) REFERENCES project_objectives(id) ON DELETE SET NULL
        )
    """)
    for column, definition in [
        ("project_id", "TEXT"),
        ("objective_id", "INTEGER"),
        ("name", "TEXT NOT NULL DEFAULT ''"),
        ("kind", "TEXT NOT NULL DEFAULT 'quantity'"),
        ("unit", "TEXT NOT NULL DEFAULT ''"),
        ("baseline", "REAL NOT NULL DEFAULT 0"),
        ("target", "REAL NOT NULL DEFAULT 100"),
        ("current_value", "REAL NOT NULL DEFAULT 0"),
        ("expected_value", "REAL"),
        ("weight", "REAL NOT NULL DEFAULT 100"),
        ("evidence_required", "INTEGER NOT NULL DEFAULT 0"),
        ("last_updated", "TEXT"),
    ]:
        add_column_if_missing(conn, "project_metrics", column, definition)

    # Backfill universal aliases when this database was created by the older v1 schema.
    metric_columns_now = {
        row["name"] for row in conn.execute("PRAGMA table_info(project_metrics)").fetchall()
    }
    if "metric_kind" in metric_columns_now and "kind" in metric_columns_now:
        conn.execute("""
            UPDATE project_metrics
            SET kind = metric_kind
            WHERE metric_kind IS NOT NULL AND TRIM(metric_kind) <> ''
        """)
    if "target_value" in metric_columns_now and "target" in metric_columns_now:
        conn.execute("""
            UPDATE project_metrics
            SET target = target_value
            WHERE target_value IS NOT NULL
        """)

    conn.execute("""
        CREATE TABLE IF NOT EXISTS project_evidence (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            project_id TEXT NOT NULL,
            metric_id INTEGER,
            evidence_type TEXT NOT NULL DEFAULT 'manual',
            source_name TEXT NOT NULL DEFAULT '',
            reference TEXT NOT NULL DEFAULT '',
            extracted_value REAL,
            confidence REAL NOT NULL DEFAULT 0,
            notes TEXT NOT NULL DEFAULT '',
            created_at TEXT NOT NULL,
            FOREIGN KEY(project_id) REFERENCES projects(id) ON DELETE CASCADE,
            FOREIGN KEY(metric_id) REFERENCES project_metrics(id) ON DELETE SET NULL
        )
    """)
    for column, definition in [
        ("project_id", "TEXT"),
        ("metric_id", "INTEGER"),
        ("evidence_type", "TEXT NOT NULL DEFAULT 'manual'"),
        ("source_name", "TEXT NOT NULL DEFAULT ''"),
        ("reference", "TEXT NOT NULL DEFAULT ''"),
        ("extracted_value", "REAL"),
        ("confidence", "REAL NOT NULL DEFAULT 0"),
        ("notes", "TEXT NOT NULL DEFAULT ''"),
        ("created_at", "TEXT NOT NULL DEFAULT ''"),
    ]:
        add_column_if_missing(conn, "project_evidence", column, definition)

    # Evidence lifecycle metadata. observed_at is the date the source describes;
    # created_at is only when it entered AIPA. applied_at means a value was reviewed/applied.
    add_column_if_missing(conn, "project_evidence", "observed_at", "TEXT")
    add_column_if_missing(conn, "project_evidence", "applied_at", "TEXT")
    add_column_if_missing(conn, "project_evidence", "is_current", "INTEGER NOT NULL DEFAULT 0")
    add_column_if_missing(conn, "project_evidence", "extracted_unit", "TEXT NOT NULL DEFAULT ''")

    conn.execute("""
        CREATE TABLE IF NOT EXISTS project_updates (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            project_id TEXT NOT NULL,
            message TEXT NOT NULL,
            source TEXT NOT NULL DEFAULT 'manual',
            created_at TEXT NOT NULL
        )
    """)
    for column, definition in [
        ("project_id", "TEXT"),
        ("message", "TEXT NOT NULL DEFAULT ''"),
        ("source", "TEXT NOT NULL DEFAULT 'manual'"),
        ("created_at", "TEXT NOT NULL DEFAULT ''"),
    ]:
        add_column_if_missing(conn, "project_updates", column, definition)


    count = conn.execute("SELECT COUNT(*) FROM projects").fetchone()[0]
    if count == 0:
        initial_projects = [
            ("PRJ-001", "Hyderabad Regional Road Development", "Roads", "Telangana", "Hyderabad", 1000, 620, 48, 67, 0, "Low", "Low", "In Progress"),
            ("PRJ-002", "Metro Rail Expansion Phase II", "Urban Transport", "Telangana", "Hyderabad", 2500, 1450, 62, 68, 0, "Low", "Low", "In Progress"),
            ("PRJ-003", "National Highway Widening Project", "Roads", "Maharashtra", "Nagpur", 1800, 900, 73, 70, 0, "Low", "Low", "In Progress"),
            ("PRJ-004", "Rural Water Supply Programme", "Water", "Karnataka", "Bengaluru", 750, 510, 51, 64, 0, "Low", "Low", "At Risk"),
            ("PRJ-005", "Solar Power Infrastructure Project", "Energy", "Rajasthan", "Jodhpur", 1300, 680, 69, 71, 0, "Low", "Low", "On Track"),
        ]
        conn.executemany("""
            INSERT INTO projects (
                id,name,sector,state,location,approved_cost,current_expenditure,
                physical_progress,expected_progress,risk_score,delay_risk,cost_risk,status
            ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)
        """, initial_projects)

    # Existing projects are legacy records. Give them a project type without
    # forcing them to invent objectives/metrics.
    rows = conn.execute("SELECT id, sector, project_type FROM projects").fetchall()
    for row in rows:
        if not row["project_type"] or row["project_type"] == "Infrastructure":
            inferred = normalize_project_type(None, row["sector"])
            conn.execute(
                "UPDATE projects SET project_type = ? WHERE id = ?",
                (inferred, row["id"]),
            )

    # Legacy evidence created before the evidence-lifecycle fields existed may
    # already contain a reviewed/extracted value but have no applied_at timestamp.
    # Preserve those observations for Phase 7 history without changing which
    # evidence is currently authoritative for the metric.
    conn.execute("""
        UPDATE project_evidence
        SET applied_at = COALESCE(NULLIF(applied_at, ''), created_at)
        WHERE extracted_value IS NOT NULL
          AND (applied_at IS NULL OR TRIM(applied_at) = '')
    """)
    conn.commit()
    conn.close()


init_db()


# ============================================================
# UNIVERSAL PROGRESS ENGINE
# ============================================================

def metric_progress(metric) -> float:
    baseline = float(metric["baseline"] or 0)
    target = float(metric["target"] or 0)
    current = float(metric["current_value"] or 0)

    denominator = target - baseline
    if denominator <= 0:
        return 100.0 if current >= target else 0.0

    return max(0.0, min(100.0, ((current - baseline) / denominator) * 100.0))


def calculate_universal_progress(conn, project_id: str) -> dict:
    objectives = conn.execute("""
        SELECT * FROM project_objectives
        WHERE project_id = ?
        ORDER BY sort_order, id
    """, (project_id,)).fetchall()

    metrics = conn.execute("""
        SELECT * FROM project_metrics
        WHERE project_id = ?
        ORDER BY id
    """, (project_id,)).fetchall()

    if not metrics:
        row = conn.execute("""
            SELECT physical_progress, expected_progress
            FROM projects WHERE id = ?
        """, (project_id,)).fetchone()
        if not row:
            return {"progress": 0, "expected_progress": 0, "confidence": None}
        return {
            "progress": round(float(row["physical_progress"]), 2),
            "expected_progress": round(float(row["expected_progress"]), 2),
            "confidence": None,
        }

    objective_scores = []
    ungrouped = []

    for objective in objectives:
        linked = [m for m in metrics if m["objective_id"] == objective["id"]]
        if not linked:
            continue
        total_weight = sum(max(0.0, float(m["weight"] or 0)) for m in linked) or 1.0
        score = sum(metric_progress(m) * max(0.0, float(m["weight"] or 0)) for m in linked) / total_weight
        objective_scores.append((score, max(0.0, float(objective["weight"] or 0))))

    for metric in metrics:
        if metric["objective_id"] is None:
            ungrouped.append(metric)

    if objective_scores:
        groups = list(objective_scores)
        if ungrouped:
            ungrouped_weight = sum(max(0.0, float(m["weight"] or 0)) for m in ungrouped) or 1.0
            ungrouped_score = sum(
                metric_progress(m) * max(0.0, float(m["weight"] or 0))
                for m in ungrouped
            ) / ungrouped_weight
            groups.append((ungrouped_score, ungrouped_weight))
        total_group_weight = sum(weight for _, weight in groups) or 1.0
        progress = sum(score * weight for score, weight in groups) / total_group_weight
    else:
        total_weight = sum(max(0.0, float(m["weight"] or 0)) for m in metrics) or 1.0
        progress = sum(
            metric_progress(m) * max(0.0, float(m["weight"] or 0))
            for m in metrics
        ) / total_weight

    expected_metrics = [m for m in metrics if m["expected_value"] is not None]
    if expected_metrics:
        expected_weight_total = sum(max(0.0, float(m["weight"] or 0)) for m in expected_metrics) or 1.0
        expected_progress = sum(
            max(0.0, min(100.0, (
                ((float(m["expected_value"]) - float(m["baseline"] or 0)) /
                 max(0.000001, float(m["target"] or 0) - float(m["baseline"] or 0))) * 100
            ))) * max(0.0, float(m["weight"] or 0))
            for m in expected_metrics
        ) / expected_weight_total
    else:
        row = conn.execute("SELECT expected_progress FROM projects WHERE id = ?", (project_id,)).fetchone()
        expected_progress = float(row["expected_progress"]) if row else progress

    evidence_count = conn.execute("""
        SELECT COUNT(*) FROM project_evidence
        WHERE project_id = ? AND extracted_value IS NOT NULL AND applied_at IS NOT NULL
    """, (project_id,)).fetchone()[0]
    confidence = None
    if evidence_count:
        confidence = conn.execute("""
            SELECT AVG(confidence) FROM project_evidence
            WHERE project_id = ? AND extracted_value IS NOT NULL AND applied_at IS NOT NULL
        """, (project_id,)).fetchone()[0]

    return {
        "progress": round(max(0, min(100, progress)), 2),
        "expected_progress": round(max(0, min(100, expected_progress)), 2),
        "confidence": round(float(confidence), 1) if confidence is not None else None,
    }


def sync_project_progress(conn, project_id: str):
    universal = calculate_universal_progress(conn, project_id)
    row = conn.execute("""
        SELECT approved_cost, current_expenditure
        FROM projects WHERE id = ?
    """, (project_id,)).fetchone()
    if not row:
        return universal

    risk = calculate_risk(
        row["approved_cost"],
        row["current_expenditure"],
        universal["progress"],
        universal["expected_progress"],
    )

    conn.execute("""
        UPDATE projects
        SET physical_progress = ?,
            expected_progress = ?,
            risk_score = ?,
            delay_risk = ?,
            cost_risk = ?
        WHERE id = ?
    """, (
        universal["progress"],
        universal["expected_progress"],
        risk["risk_score"],
        risk["delay_risk"],
        risk["cost_risk"],
        project_id,
    ))
    return universal


# ============================================================
# SERIALIZATION
# ============================================================

def serialize_project(conn, row, include_details=True):
    project_id = row["id"]
    universal = calculate_universal_progress(conn, project_id)

    objectives_out = []
    metrics_out = []
    if include_details:
        objectives = conn.execute("""
            SELECT * FROM project_objectives
            WHERE project_id = ?
            ORDER BY sort_order, id
        """, (project_id,)).fetchall()

        for obj in objectives:
            linked_metrics = conn.execute("""
                SELECT * FROM project_metrics
                WHERE objective_id = ?
                ORDER BY id
            """, (obj["id"],)).fetchall()
            objectives_out.append({
                "id": obj["id"],
                "title": obj["title"],
                "description": obj["description"],
                "weight": obj["weight"],
                "metrics": [
                    {
                        "id": m["id"],
                        "name": m["name"],
                        "kind": m["kind"],
                        "unit": m["unit"],
                        "baseline": m["baseline"],
                        "target": m["target"],
                        "current_value": m["current_value"],
                        "expected_value": m["expected_value"],
                        "weight": m["weight"],
                        "progress": round(metric_progress(m), 2),
                        "evidence_required": bool(m["evidence_required"]),
                        "last_updated": m["last_updated"],
                    }
                    for m in linked_metrics
                ],
            })

        metrics_out = [
            {
                "id": m["id"],
                "objective_id": m["objective_id"],
                "name": m["name"],
                "kind": m["kind"],
                "unit": m["unit"],
                "baseline": m["baseline"],
                "target": m["target"],
                "current_value": m["current_value"],
                "expected_value": m["expected_value"],
                "weight": m["weight"],
                "progress": round(metric_progress(m), 2),
                "evidence_required": bool(m["evidence_required"]),
                "last_updated": m["last_updated"],
            }
            for m in conn.execute(
                "SELECT * FROM project_metrics WHERE project_id = ? ORDER BY id",
                (project_id,),
            ).fetchall()
        ]

    return {
        "id": row["id"],
        "name": row["name"],
        "sector": row["sector"],
        "state": row["state"],
        "location": row["location"],
        "approved_cost": row["approved_cost"],
        "current_expenditure": row["current_expenditure"],
        "physical_progress": row["physical_progress"],
        "expected_progress": row["expected_progress"],
        "risk_score": row["risk_score"],
        "delay_risk": row["delay_risk"],
        "cost_risk": row["cost_risk"],
        "status": row["status"],
        "project_type": row["project_type"],
        "description": row["description"],
        "start_date": row["start_date"],
        "target_date": row["target_date"],
        "progress_model": row["progress_model"],
        "overall_progress": universal["progress"],
        "expected_overall_progress": universal["expected_progress"],
        "measurement_confidence": universal["confidence"],
        "objectives": objectives_out,
        "metrics": metrics_out,
    }


# ============================================================
# PYDANTIC MODELS
# ============================================================

MetricKind = Literal[
    "quantity", "percentage", "currency", "beneficiaries",
    "area", "capacity", "count", "milestone", "custom"
]
EvidenceType = Literal[
    "manual", "csv", "excel", "api", "report",
    "document", "field_update", "survey"
]


class MetricCreate(BaseModel):
    name: str = Field(..., min_length=2, max_length=150)
    kind: MetricKind = "quantity"
    unit: str = ""
    baseline: float = 0
    target: float = Field(..., gt=0)
    current_value: float = 0
    expected_value: Optional[float] = None
    weight: float = Field(100, gt=0)
    evidence_required: bool = False


class ObjectiveCreate(BaseModel):
    title: str = Field(..., min_length=2, max_length=150)
    description: str = ""
    weight: float = Field(100, gt=0)
    metrics: list[MetricCreate] = Field(default_factory=list)


class ProjectCreate(BaseModel):
    name: str = Field(..., min_length=2, max_length=150)
    sector: str = Field(..., min_length=2, max_length=100)
    state: str = Field(..., min_length=2, max_length=100)
    location: str = Field(..., min_length=2, max_length=150)
    approved_cost: float = Field(0, ge=0)
    current_expenditure: float = Field(0, ge=0)
    physical_progress: float = Field(0, ge=0, le=100)
    expected_progress: float = Field(0, ge=0, le=100)
    status: Literal["On Track", "In Progress", "At Risk", "Delayed", "Completed"] = "In Progress"
    project_type: Optional[str] = None
    description: str = ""
    start_date: Optional[str] = None
    target_date: Optional[str] = None
    progress_model: Literal["universal", "legacy"] = "universal"
    objectives: list[ObjectiveCreate] = Field(default_factory=list)


class MetricValueUpdate(BaseModel):
    current_value: float
    expected_value: Optional[float] = None
    evidence_type: EvidenceType = "manual"
    source_name: str = ""
    reference: str = ""
    confidence: float = Field(50, ge=0, le=100)
    notes: str = ""


class EvidenceCreate(BaseModel):
    metric_id: Optional[str] = None
    evidence_type: EvidenceType = "manual"
    observed_at: Optional[str] = None
    source_name: str = ""
    reference: str = ""
    extracted_value: Optional[float] = None
    confidence: float = Field(50, ge=0, le=100)
    notes: str = ""


class EvidenceExtractRequest(BaseModel):
    metric_id: Optional[str] = None
    apply: bool = False
    selected_value: Optional[float] = None


class EvidenceRecommendationRequest(BaseModel):
    metric_id: Optional[str] = None


# ============================================================
# HELPERS
# ============================================================

def project_exists(conn, project_id: str):
    row = conn.execute("SELECT * FROM projects WHERE id = ?", (project_id,)).fetchone()
    if row is None:
        raise HTTPException(status_code=404, detail=f"Project {project_id} not found")
    return row


def metric_belongs_to_project(conn, project_id: str, metric_id: str):
    columns = table_columns(conn, "project_metrics")
    if "project_id" in columns:
        return conn.execute(
            "SELECT * FROM project_metrics WHERE id = ? AND project_id = ?",
            (metric_id, project_id),
        ).fetchone()
    return conn.execute("""
        SELECT m.*
        FROM project_metrics m
        JOIN project_objectives o ON o.id = m.objective_id
        WHERE m.id = ? AND o.project_id = ?
    """, (metric_id, project_id)).fetchone()


def table_columns(conn, table: str) -> dict[str, dict]:
    """Return SQLite column metadata keyed by column name."""
    return {
        row["name"]: dict(row)
        for row in conn.execute(f"PRAGMA table_info({table})").fetchall()
    }


def insert_row_compatible(conn, table: str, values: dict) -> int | str | None:
    """Insert only columns that exist, supporting both legacy and current schemas."""
    columns = table_columns(conn, table)
    usable = {key: value for key, value in values.items() if key in columns}
    if not usable:
        raise RuntimeError(f"No compatible columns available for {table}")

    cols = list(usable)
    placeholders = ",".join("?" for _ in cols)
    cursor = conn.execute(
        f"INSERT INTO {table} ({','.join(cols)}) VALUES ({placeholders})",
        [usable[col] for col in cols],
    )

    id_meta = columns.get("id")
    if id_meta:
        id_type = str(id_meta.get("type") or "").upper()
        # INTEGER PRIMARY KEY aliases SQLite's rowid.
        if id_meta.get("pk") == 1 and "INT" in id_type:
            return cursor.lastrowid
        # Legacy schemas used TEXT ids, supplied explicitly below.
    return usable.get("id", cursor.lastrowid)


def insert_objectives_and_metrics(conn, project_id: str, objectives: list[ObjectiveCreate]):
    """Insert universal data while remaining compatible with the old v1 schema."""
    objective_columns = table_columns(conn, "project_objectives")
    metric_columns = table_columns(conn, "project_metrics")

    for obj_index, objective in enumerate(objectives):
        objective_id_value = uuid4().hex

        objective_values = {
            "id": objective_id_value,
            "project_id": project_id,
            "title": objective.title.strip(),
            "name": objective.title.strip(),  # legacy v1 alias
            "description": objective.description.strip(),
            "weight": objective.weight,
            "sort_order": obj_index,
            "created_at": now_iso(),
        }

        # For the current INTEGER PRIMARY KEY schema, do not force a text id.
        if objective_columns.get("id"):
            id_meta = objective_columns["id"]
            id_type = str(id_meta.get("type") or "").upper()
            if id_meta.get("pk") == 1 and "INT" in id_type:
                objective_values.pop("id", None)

        objective_id = insert_row_compatible(conn, "project_objectives", objective_values)

        for metric in objective.metrics:
            metric_id_value = uuid4().hex
            metric_values = {
                "id": metric_id_value,
                "project_id": project_id,
                "objective_id": objective_id,
                "name": metric.name.strip(),
                "kind": metric.kind,
                "metric_kind": metric.kind,  # legacy v1 alias
                "unit": metric.unit.strip(),
                "baseline": metric.baseline,
                "target": metric.target,
                "target_value": metric.target,  # legacy v1 alias
                "current_value": metric.current_value,
                "expected_value": metric.expected_value,
                "weight": metric.weight,
                "evidence_required": int(metric.evidence_required),
                "last_updated": now_iso(),
                "source_type": "manual",  # legacy v1 alias
                "source_reference": "",   # legacy v1 alias
            }

            if metric_columns.get("id"):
                id_meta = metric_columns["id"]
                id_type = str(id_meta.get("type") or "").upper()
                if id_meta.get("pk") == 1 and "INT" in id_type:
                    metric_values.pop("id", None)

            insert_row_compatible(conn, "project_metrics", metric_values)


# ============================================================
# ROUTES
# IMPORTANT: static routes are before /{project_id}.
# ============================================================

@router.get("/alerts")
def get_early_warnings():
    conn = get_db()
    try:
        rows = conn.execute("SELECT * FROM projects ORDER BY CAST(SUBSTR(id,5) AS INTEGER)").fetchall()
        warnings = []
        for row in rows:
            risk = calculate_risk(
                row["approved_cost"], row["current_expenditure"],
                row["physical_progress"], row["expected_progress"]
            )

            used_alert_ids = {item["id"] for item in warnings}

            def add(code, category, level, title, message, value=None, threshold=None, key_suffix=None):
                # Alert IDs must be unique even when one project has multiple
                # metrics producing the same signal code. Older records or
                # backend versions may omit metric_id, so enforce uniqueness
                # at the final serialization boundary as well.
                suffix = f"-{key_suffix}" if key_suffix is not None else ""
                base_id = f"ALT-{row['id']}-{code}{suffix}"
                alert_id = base_id
                duplicate_number = 2
                while alert_id in used_alert_ids:
                    alert_id = f"{base_id}-{duplicate_number}"
                    duplicate_number += 1
                used_alert_ids.add(alert_id)
                warnings.append({
                    "id": alert_id,
                    "project_id": row["id"],
                    "project_name": row["name"],
                    "category": category,
                    "level": level,
                    "title": title,
                    "message": message,
                    "value": round(value, 2) if value is not None else None,
                    "threshold": threshold,
                    "status": "Active",
                })

            if risk["risk_score"] >= 80:
                add("RISK", "Overall Risk", "Critical", "Critical project risk detected",
                    f"{row['name']} has an overall risk score of {risk['risk_score']}/100.", risk["risk_score"], 80)
            elif risk["risk_score"] >= 65:
                add("RISK", "Overall Risk", "High", "High project risk detected",
                    f"{row['name']} has an overall risk score of {risk['risk_score']}/100.", risk["risk_score"], 65)

            if risk["progress_gap"] >= 20:
                add("DELAY", "Schedule", "Critical", "Critical schedule deviation",
                    f"{row['name']} is {risk['progress_gap']:.1f}% behind expected progress.", risk["progress_gap"], 20)
            elif risk["progress_gap"] >= 16.25:
                add("DELAY", "Schedule", "High", "High schedule deviation",
                    f"{row['name']} is {risk['progress_gap']:.1f}% behind expected progress.", risk["progress_gap"], 16.25)
            elif risk["progress_gap"] >= 5:
                level = "High" if risk["delay_score"] >= 65 else "Medium"
                add("DELAY", "Schedule", level, "Schedule delay warning",
                    f"{row['name']} is {risk['progress_gap']:.1f}% behind expected progress.", risk["progress_gap"], 5)

            if risk["cost_gap"] >= 20:
                add("COST", "Cost", "Critical", "Critical cost deviation",
                    f"{row['name']} has used {risk['expenditure_percent']:.1f}% of approved cost while progress is {row['physical_progress']:.1f}%.",
                    risk["cost_gap"], 20)
            elif risk["cost_gap"] >= 16.25:
                add("COST", "Cost", "High", "High cost deviation",
                    f"{row['name']} is spending {risk['cost_gap']:.1f} percentage points ahead of progress.",
                    risk["cost_gap"], 16.25)
            elif risk["cost_gap"] >= 5:
                level = "High" if risk["cost_score"] >= 65 else "Medium"
                add("COST", "Cost", level, "Cost warning",
                    f"{row['name']} has used {risk['expenditure_percent']:.1f}% of approved cost against {row['physical_progress']:.1f}% progress.",
                    risk["cost_gap"], 5)

            # Add evidence/metric intelligence to portfolio alerts. Risk alerts above
            # remain the primary financial/schedule warnings; these additional signals
            # surface measurement-quality problems that the risk formula cannot see.
            intelligence = build_project_intelligence(conn, row["id"])
            for signal in intelligence["signals"]:
                if signal["code"] in {"MISSING_EVIDENCE", "LOW_CONFIDENCE", "METRIC_GAP", "EVIDENCE_CONFLICT", "EVIDENCE_VARIANCE", "STALE_EVIDENCE"}:
                    add(
                        signal["code"],
                        "Evidence" if signal["code"].startswith("EVIDENCE") or signal["code"] in {"MISSING_EVIDENCE", "LOW_CONFIDENCE", "STALE_EVIDENCE"} else "Metric",
                        signal["level"],
                        signal["title"],
                        f"{row['name']}: {signal['message']}",
                        signal.get("value"),
                        signal.get("threshold"),
                        signal.get("metric_id"),
                    )

        order = {"Critical": 0, "High": 1, "Medium": 2, "Low": 3}
        warnings.sort(key=lambda item: (order[item["level"]], item["project_id"], item["category"]))
        counts = {key.lower(): sum(w["level"] == key for w in warnings) for key in order}
        return {"total": len(warnings), "active": len(warnings), "counts": counts, "alerts": warnings}
    finally:
        conn.close()


@router.get("/measurement-summary")
def measurement_summary():
    conn = get_db()
    try:
        rows = conn.execute("SELECT * FROM projects").fetchall()
        total_metrics = conn.execute("SELECT COUNT(*) FROM project_metrics").fetchone()[0]
        total_evidence = conn.execute("SELECT COUNT(*) FROM project_evidence").fetchone()[0]
        universal_projects = sum(1 for row in rows if row["progress_model"] == "universal")
        with_evidence = conn.execute("""
            SELECT COUNT(DISTINCT project_id) FROM project_evidence
        """).fetchone()[0]
        return {
            "projects": len(rows),
            "universal_projects": universal_projects,
            "metrics": total_metrics,
            "evidence_records": total_evidence,
            "projects_with_evidence": with_evidence,
        }
    finally:
        conn.close()


@router.get("/types")
def get_project_types():
    return {
        "types": [
            {
                "name": name,
                "description": {
                    "Infrastructure": "Physical assets, construction and capital works.",
                    "Healthcare": "Facilities, services, coverage and public health delivery.",
                    "Education": "Schools, services, learners and education outcomes.",
                    "Welfare": "Benefits, social programmes and beneficiary reach.",
                    "Digital": "Platforms, modules, services and digital adoption.",
                    "Environment": "Restoration, conservation and ecological targets.",
                    "Agriculture": "Farmers, land, irrigation and agricultural delivery.",
                    "Housing": "Housing delivery, completion and occupancy.",
                    "Water": "Water connections, sanitation, coverage and capacity.",
                    "Energy": "Generation, installed capacity and connections.",
                    "Transport": "Rail, metro, transit, roads and mobility services.",
                    "Administrative": "Government services, processes and adoption.",
                    "Other": "Custom public-sector programmes.",
                }[name],
                "metric_templates": [
                    {"objective": o, "name": n, "kind": k, "weight": w, "unit": u}
                    for o, n, k, w, u in METRIC_TEMPLATES[name]
                ],
            }
            for name in PROJECT_TYPES
        ]
    }


@router.get("/")
def get_projects():
    conn = get_db()
    try:
        rows = conn.execute("""
            SELECT * FROM projects
            ORDER BY CAST(SUBSTR(id,5) AS INTEGER)
        """).fetchall()
        projects = []
        for row in rows:
            if row["progress_model"] == "universal":
                sync_project_progress(conn, row["id"])
        conn.commit()
        rows = conn.execute("""
            SELECT * FROM projects
            ORDER BY CAST(SUBSTR(id,5) AS INTEGER)
        """).fetchall()
        projects = [serialize_project(conn, row, include_details=False) for row in rows]
        return {"total": len(projects), "projects": projects}
    finally:
        conn.close()


@router.get("/{project_id}/alerts")
def get_project_early_warnings(project_id: str):
    conn = get_db()
    try:
        row = project_exists(conn, project_id)
        # Reuse portfolio warning logic for one project.
        risk = calculate_risk(row["approved_cost"], row["current_expenditure"],
                              row["physical_progress"], row["expected_progress"])
        alerts = []
        if risk["risk_score"] >= 65:
            alerts.append({
                "id": f"ALT-{project_id}-RISK", "project_id": project_id,
                "project_name": row["name"], "category": "Overall Risk",
                "level": "Critical" if risk["risk_score"] >= 80 else "High",
                "title": "Project risk requires review",
                "message": f"Overall risk score is {risk['risk_score']}/100.",
                "status": "Active"
            })
        return {"project_id": project_id, "total": len(alerts), "alerts": alerts}
    finally:
        conn.close()


def safe_evidence_filename(filename: str) -> str:
    """Keep uploaded filenames filesystem-safe while preserving the extension."""
    name = Path(filename or "evidence").name.strip() or "evidence"
    return name.replace(" ", "_")


def evidence_file_reference(project_id: str, token: str) -> str:
    return f"/api/projects/{project_id}/evidence/file/{token}"


def evidence_storage_path(project_id: str, row) -> Path:
    reference = str(row["reference"] or "")
    try:
        if not reference and "source_reference" in row.keys():
            reference = str(row["source_reference"] or "")
    except Exception:
        pass
    token = reference.rstrip("/").split("/")[-1]
    filename = safe_evidence_filename(str(row["source_name"] or "evidence"))
    if not token or token.startswith("http"):
        raise HTTPException(status_code=400, detail="This evidence record does not contain a stored device file")
    path = UPLOAD_DIR / project_id / f"{token}_{filename}"
    if not path.exists() or not path.is_file():
        raise HTTPException(status_code=404, detail="Evidence file is missing from storage")
    return path


def _number_candidates(
    text: str,
    metric_name: str = "",
    unit: str = "",
    target: Optional[float] = None,
    baseline: Optional[float] = None,
    metric_kind: str = "",
) -> list[dict]:
    """Find numeric candidates and score them against metric context.

    This is deliberately conservative: targets, percentages, calculated
    differences, remaining quantities and cross-reference values are weaker
    evidence than a value explicitly reported for the requested metric.
    """
    clean = re.sub(r"\s+", " ", text or "").strip()
    candidates: list[dict] = []
    pattern = r"[-+]?\d+(?:[,.]\d+)*(?:\s*(?:%|million|billion|thousand|km|m|mw|gw|ha|hectares|units|people|students|houses|connections|households|milestones))?"
    numbers = list(re.finditer(pattern, clean, re.IGNORECASE))
    metric_terms = [t.lower() for t in re.findall(r"[a-zA-Z0-9]+", metric_name or "") if len(t) > 2]
    normalized_unit = (unit or "").strip().lower()

    for match in numbers:
        raw = match.group(0).strip()
        numeric = re.sub(r"[^0-9.+-]", "", raw.replace(",", ""))
        try:
            value = float(numeric)
        except ValueError:
            continue

        start = max(0, match.start() - 180)
        end = min(len(clean), match.end() + 180)
        context = clean[start:end]
        lower_context = context.lower()
        overlap = sum(1 for term in metric_terms if term in lower_context)
        unit_match = bool(normalized_unit and normalized_unit in lower_context)

        confidence = 38 + min(36, overlap * 12) + (12 if unit_match else 0)
        reasons: list[str] = []
        if overlap:
            reasons.append("metric context match")
        if unit_match:
            reasons.append("unit match")

        # Strongly penalize values that are normally metadata/calculations,
        # rather than the measured value being reported.
        negative_terms = (
            "target", "contracted target", "remaining", "previous month",
            "calculated", "cross-reference", "cross reference", "variance",
            "difference", "gap", "rate", "% of", "percentage", "utilisation",
            "utilization", "balance available", "implied", "added during",
        )
        if any(term in lower_context for term in negative_terms):
            confidence -= 30
            reasons.append("calculated/target context")

        # Negative numbers are generally deltas/variances for project metrics,
        # not measured completion values.
        if value < 0 and metric_kind not in {"currency"}:
            confidence -= 35
            reasons.append("negative/difference value")

        if target is not None and target > 0:
            distance = abs(value - target) / max(abs(target), 1.0)
            if distance <= 0.10:
                # Being near the target is not evidence by itself. Keep only a
                # small signal rather than rewarding the target value heavily.
                confidence += 2
                reasons.append("near target")
            elif distance <= 0.25:
                confidence += 1
                reasons.append("near target range")

        if baseline is not None and target is not None and target > baseline and baseline <= value <= target:
            confidence += 2
            reasons.append("within baseline-target range")

        if metric_kind == "percentage" and 0 <= value <= 100:
            confidence += 4
            reasons.append("valid percentage range")

        candidates.append({
            "value": value,
            "confidence": min(98, max(1, confidence)),
            "context": context,
            "reason": ", ".join(reasons) or "numeric value found",
        })

    best_by_value: dict[float, dict] = {}
    for candidate in candidates:
        current = best_by_value.get(candidate["value"])
        if current is None or candidate["confidence"] > current["confidence"]:
            best_by_value[candidate["value"]] = candidate

    candidates = list(best_by_value.values())
    candidates.sort(key=lambda item: (-item["confidence"], item["value"]))
    return candidates[:10]


def _parse_reported_value(raw_value) -> Optional[float]:
    """Parse a human-facing Evidence Association reported value."""
    if raw_value is None:
        return None
    text = str(raw_value).strip().replace(",", "")
    match = re.search(r"[-+]?\d+(?:\.\d+)?", text)
    if not match:
        return None
    try:
        return float(match.group(0))
    except ValueError:
        return None


def _metric_name_match(source_name: str, metric_name: str) -> float:
    """Return a simple explainable similarity score between two metric labels."""
    source = {t.lower() for t in re.findall(r"[a-zA-Z0-9]+", source_name or "") if len(t) > 2}
    target = {t.lower() for t in re.findall(r"[a-zA-Z0-9]+", metric_name or "") if len(t) > 2}
    if not source or not target:
        return 0.0
    overlap = len(source & target)
    return overlap / max(len(target), 1)


def _excel_evidence_association_candidates(
    path: Path,
    metric_name: str,
    unit: str,
) -> list[dict]:
    """Prefer an explicit Evidence Association sheet when a workbook has one.

    These workbooks are evidence packages, not arbitrary numeric spreadsheets.
    If the source explicitly states `Associated indicator` and `Reported value`,
    that claim should outrank unrelated targets, percentages and calculations.
    """
    try:
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True)
        candidates: list[dict] = []

        for sheet in workbook.worksheets:
            title = str(sheet.title or "").strip().lower()
            if "evidence association" not in title and "evidence" not in title:
                continue

            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                continue

            header_idx = None
            header_map = {}
            for idx, row in enumerate(rows[:12]):
                normalized = [str(v or "").strip().lower() for v in row]
                if "associated indicator" in normalized and "reported value" in normalized:
                    header_idx = idx
                    header_map = {value: pos for pos, value in enumerate(normalized)}
                    break

            if header_idx is None:
                continue

            for row in rows[header_idx + 1:]:
                def cell(name: str):
                    pos = header_map.get(name)
                    return row[pos] if pos is not None and pos < len(row) else None

                associated = str(cell("associated indicator") or "").strip()
                reported_raw = cell("reported value")
                reported = _parse_reported_value(reported_raw)
                if not associated or reported is None:
                    continue

                similarity = _metric_name_match(associated, metric_name)
                if similarity < 0.34:
                    continue

                confidence_raw = cell("confidence")
                try:
                    source_confidence = float(confidence_raw) * 100 if float(confidence_raw) <= 1 else float(confidence_raw)
                except (TypeError, ValueError):
                    source_confidence = 70.0

                evidence_type = str(cell("evidence type") or "document").strip()
                source_doc = str(cell("source document") or path.name).strip()
                notes = str(cell("notes") or "").strip()
                reported_text = str(reported_raw).strip()
                unit_hint = unit or ""

                extraction_confidence = min(
                    99.0,
                    82.0 + similarity * 10.0 + (5.0 if unit_hint and unit_hint.lower() in reported_text.lower() else 0.0),
                )
                context = (
                    f"Evidence Association · {associated}: {reported_text}"
                    f" · source: {source_doc}"
                    f" · type: {evidence_type}"
                    + (f" · {notes}" if notes else "")
                )

                candidates.append({
                    "value": reported,
                    "confidence": round(extraction_confidence, 1),
                    "context": context,
                    "reason": "explicit reported value in Evidence Association",
                    "source_confidence": round(max(0.0, min(100.0, source_confidence)), 1),
                    "evidence_type": evidence_type,
                })

        workbook.close()
        candidates.sort(key=lambda item: (-item["confidence"], item["value"]))
        return candidates[:10]
    except Exception:
        return []

def extract_evidence_candidates(
    path: Path,
    metric_name: str = "",
    unit: str = "",
    target: Optional[float] = None,
    baseline: Optional[float] = None,
    metric_kind: str = "",
) -> list[dict]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        text = path.read_text(encoding="utf-8-sig", errors="ignore")
        rows = list(csv.reader(io.StringIO(text)))
        if not rows:
            return []
        headers = [str(h).strip() for h in rows[0]]
        output = []
        for row in rows[1:2000]:
            for idx, cell in enumerate(row):
                if not str(cell).strip():
                    continue
                header = headers[idx] if idx < len(headers) else ""
                if metric_name and metric_name.lower() not in f"{header} {cell}".lower() and unit and unit.lower() not in f"{header} {cell}".lower():
                    continue
                output.extend(_number_candidates(f"{header}: {cell}", metric_name, unit, target, baseline, metric_kind))
        return output[:10] or _number_candidates(text, metric_name, unit, target, baseline, metric_kind)

    if suffix in {".xlsx", ".xlsm"}:
        try:
            association_candidates = _excel_evidence_association_candidates(path, metric_name, unit)
            if association_candidates:
                return association_candidates

            from openpyxl import load_workbook
            workbook = load_workbook(path, read_only=True, data_only=True)
            output = []
            for sheet in workbook.worksheets[:10]:
                rows = sheet.iter_rows(values_only=True)
                headers = [str(v).strip() if v is not None else "" for v in next(rows, [])]
                for row in rows:
                    for idx, cell in enumerate(row):
                        if cell is None:
                            continue
                        header = headers[idx] if idx < len(headers) else ""
                        label = f"{sheet.title} · {header}"
                        if isinstance(cell, (int, float)):
                            context = f"{label}: {cell}"
                            confidence = 60
                            if metric_name and metric_name.lower() in label.lower(): confidence += 25
                            if unit and unit.lower() in label.lower(): confidence += 10
                            output.append({"value": float(cell), "confidence": min(confidence, 95), "context": context})
                        else:
                            output.extend(_number_candidates(f"{label}: {cell}", metric_name, unit, target, baseline, metric_kind))
            workbook.close()
            output.sort(key=lambda item: (-item["confidence"], item["value"]))
            return output[:10]
        except Exception as exc:
            raise HTTPException(status_code=422, detail=f"Could not read Excel evidence: {exc}")

    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
            reader = PdfReader(str(path))
            text = "\n".join((page.extract_text() or "") for page in reader.pages[:50])
            return _number_candidates(text, metric_name, unit, target, baseline, metric_kind)
        except Exception as exc:
            raise HTTPException(status_code=422, detail=f"Could not read PDF evidence: {exc}")

    if suffix in {".docx"}:
        try:
            from docx import Document
            doc = Document(str(path))
            text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
            return _number_candidates(text, metric_name, unit, target, baseline, metric_kind)
        except Exception as exc:
            raise HTTPException(status_code=422, detail=f"Could not read Word evidence: {exc}")

    if suffix in {".txt", ".md", ".json", ".xml", ".html", ".htm"}:
        text = path.read_text(encoding="utf-8-sig", errors="ignore")
        return _number_candidates(text, metric_name, unit, target, baseline, metric_kind)

    if suffix in {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tiff"}:
        try:
            from PIL import Image
            import pytesseract
            text = pytesseract.image_to_string(Image.open(path))
            return _number_candidates(text, metric_name, unit, target, baseline, metric_kind)
        except Exception as exc:
            raise HTTPException(status_code=422, detail=f"Could not OCR image evidence: {exc}")

    raise HTTPException(status_code=415, detail="Automatic extraction is supported for CSV, XLSX/XLSM, PDF, DOCX, TXT, JSON, XML, HTML and common image files.")


MONTH_LOOKUP = {name.lower(): i for i, name in enumerate(calendar.month_abbr) if name}
MONTH_LOOKUP.update({name.lower(): i for i, name in enumerate(calendar.month_name) if name})

def normalize_observed_at(value: Optional[str]) -> Optional[str]:
    """Normalize a source observation date to YYYY-MM-DD when possible."""
    if not value:
        return None
    text = str(value).strip()
    if not text:
        return None
    try:
        return datetime.fromisoformat(text.replace("Z", "+00:00")).date().isoformat()
    except ValueError:
        return None

def infer_observed_at_from_name(name: str) -> Optional[str]:
    """Best-effort date inference for filenames such as Aug2026 or AugSep2026.
    Ambiguous month ranges use the later month. Explicit UI dates always win.
    """
    text = Path(name or "").stem.lower().replace("_", " ").replace("-", " ")
    year_match = re.search(r"\\b(20\\d{2})\\b", text)
    if not year_match:
        year_match = re.search(r"(20\\d{2})", text)
    if not year_match:
        return None
    year = int(year_match.group(1))
    months = []
    for token, month in MONTH_LOOKUP.items():
        if len(token) >= 3 and re.search(rf"(?<![a-z]){re.escape(token)}(?![a-z])", text):
            months.append(month)
    # Also support concatenated ranges such as AugSep2026.
    for short, month in [(calendar.month_abbr[i].lower(), i) for i in range(1, 13)]:
        if short in text and month not in months:
            months.append(month)
    if not months:
        return None
    month = max(months)
    return f"{year:04d}-{month:02d}-{calendar.monthrange(year, month)[1]:02d}"

def evidence_effective_date(row) -> str:
    observed = normalize_observed_at(row["observed_at"] if "observed_at" in row.keys() else None)
    if observed:
        return observed
    inferred = infer_observed_at_from_name(str(row["source_name"] or ""))
    if inferred:
        return inferred
    return str(row["created_at"] or "")[:10]

def reconcile_metric_from_evidence(conn, project_id: str, metric_id) -> Optional[dict]:
    """Select the newest reviewed evidence as the metric's current observation.
    Historical evidence remains available for trend/consistency but cannot roll progress backward.
    """
    rows = conn.execute("""
        SELECT * FROM project_evidence
        WHERE project_id = ? AND metric_id = ?
          AND extracted_value IS NOT NULL AND applied_at IS NOT NULL
    """, (project_id, metric_id)).fetchall()
    conn.execute("UPDATE project_evidence SET is_current = 0 WHERE project_id = ? AND metric_id = ?", (project_id, metric_id))
    if not rows:
        return None
    chosen = max(rows, key=lambda r: (evidence_effective_date(r), str(r["applied_at"] or ""), str(r["id"])))
    conn.execute("UPDATE project_evidence SET is_current = 1 WHERE project_id = ? AND CAST(id AS TEXT) = ?", (project_id, str(chosen["id"])))
    metric_columns = table_columns(conn, "project_metrics")
    updates = {"current_value": float(chosen["extracted_value"]), "last_updated": evidence_effective_date(chosen)}
    usable = {k: v for k, v in updates.items() if k in metric_columns}
    if usable:
        conn.execute(f"UPDATE project_metrics SET {', '.join(f'{k} = ?' for k in usable)} WHERE id = ?", [*usable.values(), metric_id])
    return {"evidence_id": chosen["id"], "value": float(chosen["extracted_value"]), "observed_at": evidence_effective_date(chosen)}

def _metric_history(conn, project_id: str) -> list[dict]:
    """Build an evidence-backed observation timeline and explainable pace metrics."""
    project = project_exists(conn, project_id)
    metrics = conn.execute(
        "SELECT * FROM project_metrics WHERE project_id = ? ORDER BY id",
        (project_id,),
    ).fetchall()

    history = []
    for metric in metrics:
        rows = conn.execute(
            """
            SELECT
                e.id, e.extracted_value, e.confidence, e.source_name, e.reference,
                e.observed_at, e.applied_at, e.is_current
            FROM project_evidence e
            WHERE e.project_id = ?
              AND e.metric_id = ?
              AND e.extracted_value IS NOT NULL
              AND e.applied_at IS NOT NULL
            ORDER BY COALESCE(e.observed_at, e.created_at), e.id
            """,
            (project_id, metric["id"]),
        ).fetchall()

        observations = []
        for row in rows:
            observed_at = evidence_effective_date(row)
            observations.append({
                "evidence_id": str(row["id"]),
                "observed_at": observed_at,
                "value": round(float(row["extracted_value"]), 4),
                "confidence": round(float(row["confidence"] or 0), 1),
                "source_name": str(row["source_name"] or row["reference"] or "Evidence"),
                "reference": str(row["reference"] or ""),
                "is_current": bool(row["is_current"]),
            })

        latest = observations[-1] if observations else None
        previous = observations[-2] if len(observations) >= 2 else None
        delta = None
        delta_percent = None
        days_between = None
        rate_per_day = None

        if latest and previous:
            latest_date = datetime.fromisoformat(latest["observed_at"])
            previous_date = datetime.fromisoformat(previous["observed_at"])
            days_between = max(0, (latest_date - previous_date).days)
            delta = latest["value"] - previous["value"]
            delta_percent = (
                (delta / abs(previous["value"])) * 100
                if previous["value"] != 0 else None
            )
            if days_between > 0:
                rate_per_day = delta / days_between

        first = observations[0] if observations else None
        overall_rate_per_day = None
        overall_days = None
        if first and latest and first["observed_at"] != latest["observed_at"]:
            first_date = datetime.fromisoformat(first["observed_at"])
            latest_date = datetime.fromisoformat(latest["observed_at"])
            overall_days = max(0, (latest_date - first_date).days)
            if overall_days > 0:
                overall_rate_per_day = (latest["value"] - first["value"]) / overall_days

        trend = "insufficient data"
        if len(observations) >= 2:
            if delta is not None and abs(delta) <= max(abs(previous["value"]) * 0.01, 0.001):
                trend = "stable"
            elif delta is not None and delta > 0:
                trend = "increasing"
            elif delta is not None:
                trend = "decreasing"

        target = float(metric["target"] or 0)
        current_value = float(metric["current_value"] or 0)
        target_remaining = max(0.0, target - current_value)
        required_rate_per_day = None
        days_to_target = None
        target_date_text = str(project["target_date"] or "").strip()
        if target_date_text:
            try:
                target_date = datetime.fromisoformat(target_date_text.replace("Z", "+00:00"))
                reference_date = datetime.fromisoformat((latest["observed_at"] if latest else now_iso()).replace("Z", "+00:00"))
                days_to_target = max(0, (target_date - reference_date).days)
                if days_to_target > 0:
                    required_rate_per_day = target_remaining / days_to_target
            except ValueError:
                pass

        pace_ratio = None
        pace_status = "insufficient data"
        if required_rate_per_day is not None and rate_per_day is not None:
            pace_ratio = rate_per_day / required_rate_per_day if required_rate_per_day > 0 else None
            if rate_per_day <= 0:
                pace_status = "not progressing"
            elif pace_ratio is not None and pace_ratio >= 0.9:
                pace_status = "near required pace"
            else:
                pace_status = "behind required pace"

        history.append({
            "metric_id": metric["id"],
            "metric_name": str(metric["name"] or "Metric"),
            "unit": str(metric["unit"] or ""),
            "baseline": float(metric["baseline"] or 0),
            "target": target,
            "current_value": current_value,
            "progress": round(metric_progress(metric), 2),
            "trend": trend,
            "observation_count": len(observations),
            "observations": observations,
            "latest_observed_at": latest["observed_at"] if latest else None,
            "previous_value": previous["value"] if previous else None,
            "delta": round(delta, 4) if delta is not None else None,
            "delta_percent": round(delta_percent, 2) if delta_percent is not None else None,
            "days_between": days_between,
            "rate_per_day": round(rate_per_day, 6) if rate_per_day is not None else None,
            "overall_rate_per_day": round(overall_rate_per_day, 6) if overall_rate_per_day is not None else None,
            "required_rate_per_day": round(required_rate_per_day, 6) if required_rate_per_day is not None else None,
            "target_remaining": round(target_remaining, 4),
            "days_to_target": days_to_target,
            "pace_ratio": round(pace_ratio, 3) if pace_ratio is not None else None,
            "pace_status": pace_status,
        })
    return history


@router.get("/{project_id}/metrics/history")
def get_metric_history(project_id: str):
    conn = get_db()
    try:
        return {
            "project_id": project_id,
            "metrics": _metric_history(conn, project_id),
        }
    finally:
        conn.close()


@router.get("/{project_id}/evidence/consistency")
def get_evidence_consistency(project_id: str):
    """Return explainable discrepancies between extracted values for each metric."""
    conn = get_db()
    try:
        project_exists(conn, project_id)
        rows = conn.execute("""
            SELECT e.id, e.metric_id, e.source_name, e.reference, e.extracted_value,
                   e.confidence, e.created_at, m.name AS metric_name, m.unit AS metric_unit,
                   m.current_value AS metric_current_value
            FROM project_evidence e
            JOIN project_metrics m ON m.id = e.metric_id
            WHERE e.project_id = ? AND e.extracted_value IS NOT NULL AND e.applied_at IS NOT NULL
            ORDER BY COALESCE(e.observed_at, e.created_at) DESC, e.id DESC
        """, (project_id,)).fetchall()

        grouped: dict[str, dict] = {}
        for row in rows:
            key = str(row["metric_id"])
            group = grouped.setdefault(key, {
                "metric_id": row["metric_id"],
                "metric_name": str(row["metric_name"] or "Metric"),
                "unit": str(row["metric_unit"] or ""),
                "current_value": float(row["metric_current_value"] or 0),
                "sources": [],
            })
            group["sources"].append({
                "evidence_id": row["id"],
                "source_name": str(row["source_name"] or row["reference"] or "Evidence"),
                "reference": str(row["reference"] or ""),
                "value": float(row["extracted_value"]),
                "confidence": float(row["confidence"] or 0),
                "created_at": row["created_at"],
                "observed_at": row["observed_at"] if "observed_at" in row.keys() else None,
                "is_current": bool(row["is_current"]) if "is_current" in row.keys() else False,
            })

        insights = []
        for group in grouped.values():
            values = [source["value"] for source in group["sources"]]
            minimum, maximum = min(values), max(values)
            average = sum(values) / len(values)
            spread = ((maximum - minimum) / max(abs(maximum), 1.0)) * 100
            current_gap = abs(group["current_value"] - average) / max(abs(average), 1.0) * 100

            if len(values) >= 2 and spread >= 15:
                status = "conflict"
                message = f"Evidence values differ by {spread:.1f}%. Review the source records before treating the metric as settled."
            elif len(values) >= 2 and spread >= 8:
                status = "watch"
                message = f"Evidence values show an {spread:.1f}% spread. A source review may be useful."
            elif len(values) == 1 and current_gap >= 15:
                status = "watch"
                message = f"The current metric value differs from its only extracted evidence value by {current_gap:.1f}%."
            else:
                status = "consistent"
                message = "Available extracted evidence is internally consistent."

            insights.append({
                "metric_id": group["metric_id"],
                "metric_name": group["metric_name"],
                "unit": group["unit"],
                "status": status,
                "message": message,
                "source_count": len(values),
                "minimum": round(minimum, 4),
                "maximum": round(maximum, 4),
                "average": round(average, 4),
                "spread_percent": round(spread, 2),
                "current_value": group["current_value"],
                "sources": group["sources"],
            })

        conflicts = sum(item["status"] == "conflict" for item in insights)
        watches = sum(item["status"] == "watch" for item in insights)
        return {
            "project_id": project_id,
            "insights": insights,
            "conflicts": conflicts,
            "watches": watches,
            "summary": (f"{conflicts} evidence conflict(s) and {watches} item(s) needing review." if conflicts or watches else "No evidence discrepancies detected."),
        }
    finally:
        conn.close()



# ============================================================
# EXPLAINABLE PROJECT INTELLIGENCE ENGINE
# ============================================================

def build_project_intelligence(conn, project_id: str) -> dict:
    """Build metric-level, evidence-aware signals without silently changing project data."""
    project = project_exists(conn, project_id)
    universal = calculate_universal_progress(conn, project_id)
    risk = calculate_risk(
        project["approved_cost"],
        project["current_expenditure"],
        universal["progress"],
        universal["expected_progress"],
    )

    metric_rows = conn.execute("""
        SELECT
            m.*,
            COUNT(e.id) AS evidence_count,
            AVG(e.confidence) AS evidence_confidence,
            MAX(COALESCE(e.observed_at, e.created_at)) AS latest_evidence_at
        FROM project_metrics m
        LEFT JOIN project_evidence e
            ON e.metric_id = m.id AND e.project_id = m.project_id
           AND e.extracted_value IS NOT NULL AND e.applied_at IS NOT NULL
        WHERE m.project_id = ?
        GROUP BY m.id
        ORDER BY m.id
    """, (project_id,)).fetchall()

    metric_history = _metric_history(conn, project_id)
    history_by_id = {str(item["metric_id"]): item for item in metric_history}

    signals = []
    drivers = []
    actions = []

    def add_signal(code, level, title, message, metric_id=None, value=None, threshold=None, action=""):
        signals.append({
            "id": f"{project_id}-{code}-{metric_id or 'project'}",
            "code": code,
            "level": level,
            "title": title,
            "message": message,
            "metric_id": metric_id,
            "value": round(float(value), 2) if value is not None else None,
            "threshold": threshold,
            "recommended_action": action,
        })
        if action and action not in actions:
            actions.append(action)
        drivers.append({
            "code": code,
            "level": level,
            "title": title,
            "message": message,
            "metric_id": metric_id,
        })

    if risk["progress_gap"] >= 10:
        level = "Critical" if risk["progress_gap"] >= 20 else "High" if risk["progress_gap"] >= 16.25 else "Medium"
        add_signal(
            "SCHEDULE_GAP", level,
            "Progress is behind expected delivery",
            f"Measured progress is {risk['progress_gap']:.1f} percentage points below expected progress.",
            value=risk["progress_gap"], threshold=10,
            action="Review delayed milestones, the latest field update, and the recovery plan.",
        )

    if risk["cost_gap"] >= 10:
        level = "Critical" if risk["cost_gap"] >= 20 else "High" if risk["cost_gap"] >= 16.25 else "Medium"
        add_signal(
            "COST_GAP", level,
            "Expenditure is ahead of measured progress",
            f"Cost utilisation is {risk['expenditure_percent']:.1f}% while measured progress is {universal['progress']:.1f}%.",
            value=risk["cost_gap"], threshold=10,
            action="Review recent expenditure against completed work and remaining commitments.",
        )

    for metric in metric_rows:
        baseline = float(metric["baseline"] or 0)
        target = float(metric["target"] or 0)
        expected_value = metric["expected_value"]
        metric_prog = metric_progress(metric)
        evidence_count = int(metric["evidence_count"] or 0)
        evidence_confidence = float(metric["evidence_confidence"]) if metric["evidence_confidence"] is not None else None

        if bool(metric["evidence_required"]) and evidence_count == 0:
            add_signal(
                "MISSING_EVIDENCE", "Medium",
                f"Evidence is missing for {metric['name']}",
                "This metric is marked as requiring evidence, but no evidence record is attached.",
                metric_id=metric["id"],
                action=f"Attach a current source record to the {metric['name']} metric.",
            )

        if evidence_count > 0 and evidence_confidence is not None and evidence_confidence < 60:
            add_signal(
                "LOW_CONFIDENCE", "Medium",
                f"Low evidence confidence for {metric['name']}",
                f"Attached evidence has an average confidence of {evidence_confidence:.0f}%.",
                metric_id=metric["id"], value=evidence_confidence, threshold=60,
                action=f"Review the source quality and confirm the measured value for {metric['name']}.",
            )

        if expected_value is not None and target > baseline:
            expected_metric_progress = max(
                0.0,
                min(100.0, ((float(expected_value) - baseline) / (target - baseline)) * 100.0)
            )
            metric_gap = expected_metric_progress - metric_prog
            if metric_gap >= 10:
                level = "High" if metric_gap >= 20 else "Medium"
                add_signal(
                    "METRIC_GAP", level,
                    f"{metric['name']} is behind its expected value",
                    f"Metric progress is {metric_gap:.1f} percentage points below its expected trajectory.",
                    metric_id=metric["id"], value=metric_gap, threshold=10,
                    action=f"Review the delivery path for {metric['name']} and update its latest evidence.",
                )

        if evidence_count > 0 and metric["latest_evidence_at"]:
            try:
                latest = datetime.fromisoformat(str(metric["latest_evidence_at"]).replace("Z", "+00:00"))
                age_days = max(0, (datetime.now(latest.tzinfo) - latest).days)
                if age_days >= 60:
                    add_signal(
                        "STALE_EVIDENCE", "Medium",
                        f"Evidence may be stale for {metric['name']}",
                        f"The latest attached evidence is approximately {age_days} days old.",
                        metric_id=metric["id"], value=age_days, threshold=60,
                        action=f"Request or upload a newer source for {metric['name']}.",
                    )
            except Exception:
                pass

        history_item = history_by_id.get(str(metric["id"]))
        if history_item and history_item["observation_count"] >= 2:
            trend = history_item["trend"]
            delta = history_item["delta"]
            rate = history_item["rate_per_day"]
            if trend == "decreasing":
                add_signal(
                    "METRIC_TREND_DECLINE", "High",
                    f"{metric['name']} has declined in the latest observation",
                    f"The latest reviewed observation changed by {delta:g} {metric['unit'] or 'units'} versus the previous observation.",
                    metric_id=metric["id"], value=delta,
                    action=f"Review the latest source and field status for {metric['name']}.",
                )
            elif trend == "stable":
                add_signal(
                    "METRIC_TREND_STALLED", "Medium",
                    f"{metric['name']} shows little recent movement",
                    "The latest observation is within 1% of the previous observation.",
                    metric_id=metric["id"],
                    action=f"Review whether delivery for {metric['name']} has stalled and whether the latest evidence is current.",
                )

            if history_item["pace_status"] == "behind required pace":
                add_signal(
                    "METRIC_PACE_GAP", "High",
                    f"{metric['name']} is behind the required delivery pace",
                    f"Recent delivery is about {rate * 30:.2f} {metric['unit'] or 'units'} per 30 days versus approximately {history_item['required_rate_per_day'] * 30:.2f} required.",
                    metric_id=metric["id"],
                    value=rate * 30,
                    threshold=history_item["required_rate_per_day"] * 30,
                    action=f"Review the recovery pace needed to reach the target for {metric['name']}.",
                )

    evidence_rows = conn.execute("""
        SELECT e.metric_id, e.extracted_value,
               m.name AS metric_name, m.unit AS metric_unit
        FROM project_evidence e
        JOIN project_metrics m ON m.id = e.metric_id
        WHERE e.project_id = ? AND e.extracted_value IS NOT NULL AND e.applied_at IS NOT NULL
    """, (project_id,)).fetchall()

    grouped = {}
    for row in evidence_rows:
        key = str(row["metric_id"])
        grouped.setdefault(key, {
            "metric_id": row["metric_id"],
            "metric_name": str(row["metric_name"] or "Metric"),
            "unit": str(row["metric_unit"] or ""),
            "values": [],
        })
        grouped[key]["values"].append(float(row["extracted_value"]))

    consistency_issues = 0
    consistency_warnings = 0
    for group in grouped.values():
        values = group["values"]
        if len(values) >= 2:
            spread = ((max(values) - min(values)) / max(abs(max(values)), 1.0)) * 100
            if spread >= 15:
                consistency_issues += 1
                add_signal(
                    "EVIDENCE_CONFLICT", "High",
                    f"Conflicting evidence for {group['metric_name']}",
                    f"Attached extracted values span {min(values):g} to {max(values):g} ({spread:.1f}% spread).",
                    metric_id=group["metric_id"], value=spread, threshold=15,
                    action=f"Review the source records for {group['metric_name']} before treating the metric as settled.",
                )
            elif spread >= 8:
                consistency_warnings += 1
                add_signal(
                    "EVIDENCE_VARIANCE", "Medium",
                    f"Evidence variance for {group['metric_name']}",
                    f"Attached extracted values have an {spread:.1f}% spread.",
                    metric_id=group["metric_id"], value=spread, threshold=8,
                    action=f"Check which source is the latest authoritative value for {group['metric_name']}.",
                )

    severity_order = {"Critical": 0, "High": 1, "Medium": 2, "Low": 3}
    signals.sort(key=lambda item: (severity_order[item["level"]], item["title"]))

    critical_count = sum(1 for item in signals if item["level"] == "Critical")
    high_count = sum(1 for item in signals if item["level"] == "High")
    medium_count = sum(1 for item in signals if item["level"] == "Medium")

    if critical_count:
        attention = "Critical"
    elif high_count:
        attention = "High"
    elif medium_count:
        attention = "Medium"
    else:
        attention = "Low"

    summary = (
        "No material delivery, cost, metric, or evidence-quality signals were detected from the available project data."
        if not signals
        else f"{len(signals)} signal(s) detected: {critical_count} critical, {high_count} high, {medium_count} medium."
    )

    return {
        "project_id": project_id,
        "attention_level": attention,
        "summary": summary,
        "overall_progress": universal["progress"],
        "expected_progress": universal["expected_progress"],
        "progress_gap": risk["progress_gap"],
        "risk_score": risk["risk_score"],
        "measurement_confidence": universal["confidence"],
        "evidence_conflicts": consistency_issues,
        "evidence_variances": consistency_warnings,
        "signal_count": len(signals),
        "signals": signals,
        "drivers": drivers[:5],
        "recommended_actions": actions[:6],
        "metric_history": metric_history,
        "explainability": {
            "method": "rule-based evidence and delivery analysis",
            "inputs": ["measured progress", "expected progress", "financial utilisation", "metric trajectories", "evidence confidence", "evidence consistency", "evidence recency"],
            "data_changed": False,
        },
    }


@router.get("/{project_id}/intelligence")
def get_project_intelligence(project_id: str):
    conn = get_db()
    try:
        return build_project_intelligence(conn, project_id)
    finally:
        conn.close()



def _recommend_evidence_metrics(
    conn,
    project_id: str,
    evidence_row,
    requested_metric_id: Optional[str] = None,
) -> dict:
    """Rank evidence candidates against project metrics without changing stored data."""
    path = None
    try:
        path = evidence_storage_path(project_id, evidence_row)
    except HTTPException:
        # Link/manual evidence may not have a stored file. A previously extracted
        # value can still be scored against the project's metrics.
        if evidence_row["extracted_value"] is None:
            raise
    candidates_by_metric = []

    metrics = conn.execute(
        "SELECT * FROM project_metrics WHERE project_id = ? ORDER BY id",
        (project_id,),
    ).fetchall()

    if requested_metric_id:
        metrics = [
            metric for metric in metrics
            if str(metric["id"]) == str(requested_metric_id)
        ]
        if not metrics:
            raise HTTPException(status_code=404, detail="Requested metric was not found for this project.")

    stored_value = evidence_row["extracted_value"]
    source_confidence = max(0.0, min(100.0, float(evidence_row["confidence"] or 0)))

    for metric in metrics:
        metric_name = str(metric["name"] or "")
        unit = str(metric["unit"] or "")
        target = float(metric["target"] or 0)
        baseline = float(metric["baseline"] or 0)
        metric_kind = str(metric["kind"] or "")

        if path is not None and path.exists():
            candidates = extract_evidence_candidates(
                path, metric_name, unit, target, baseline, metric_kind
            )
        elif stored_value is not None:
            candidates = [{
                "value": float(stored_value),
                "confidence": source_confidence,
                "context": f"Stored evidence value: {stored_value}",
                "reason": "existing extracted evidence value",
            }]
        else:
            candidates = []

        if not candidates:
            continue

        existing = conn.execute(
            """
            SELECT extracted_value, confidence, created_at
            FROM project_evidence
            WHERE project_id = ? AND metric_id = ?
              AND extracted_value IS NOT NULL
              AND id != ?
            ORDER BY created_at DESC
            LIMIT 25
            """,
            (project_id, metric["id"], evidence_row["id"]),
        ).fetchall()

        existing_values = [float(row["extracted_value"]) for row in existing]
        consistency_average = (
            sum(existing_values) / len(existing_values)
            if existing_values else None
        )

        for candidate in candidates:
            value = float(candidate["value"])
            extraction_confidence = max(
                0.0, min(100.0, float(candidate.get("confidence", 0)))
            )

            # An explicit Evidence Association claim is stronger than a
            # generic number found somewhere else in the workbook.
            is_explicit_report = candidate.get("reason") == "explicit reported value in Evidence Association"
            score = extraction_confidence * (0.78 if is_explicit_report else 0.52)
            reason_codes = []

            candidate_reason = str(candidate.get("reason") or "")
            if "metric context match" in candidate_reason:
                score += 14
                reason_codes.append("metric context match")
            if "unit match" in candidate_reason:
                score += 10
                reason_codes.append("unit match")
            if is_explicit_report:
                score += 12
                reason_codes.append("explicit reported value")
            if "within baseline-target range" in candidate_reason:
                score += 6
                reason_codes.append("inside target range")
            if "near target" in candidate_reason:
                score += 5
                reason_codes.append("near target")

            if metric_kind == "percentage" and not 0 <= value <= 100:
                score -= 35
                reason_codes.append("percentage range mismatch")

            current_value = float(metric["current_value"] or 0)
            trajectory_denominator = max(abs(target - baseline), 1.0)
            trajectory_gap = abs(value - current_value) / trajectory_denominator
            if trajectory_gap <= 0.10:
                score += 7
                reason_codes.append("close to current trajectory")
            elif trajectory_gap > 0.75:
                score -= 6
                reason_codes.append("far from current trajectory")

            consistency_spread = None
            if consistency_average is not None:
                consistency_spread = abs(value - consistency_average) / max(abs(consistency_average), 1.0) * 100
                if consistency_spread <= 5:
                    score += 12
                    reason_codes.append("matches existing evidence")
                elif consistency_spread <= 10:
                    score += 5
                    reason_codes.append("near existing evidence")
                elif consistency_spread >= 20:
                    score -= 18
                    reason_codes.append("conflicts with existing evidence")

            score += (source_confidence - 50) * 0.12
            if source_confidence >= 80:
                reason_codes.append("high source confidence")

            if bool(metric["evidence_required"]):
                score += 3
                reason_codes.append("evidence-required metric")

            score = max(0.0, min(99.0, score))

            explanation_parts = []
            if candidate.get("reason"):
                explanation_parts.append(str(candidate["reason"]))
            if consistency_average is not None:
                explanation_parts.append(
                    f"{consistency_spread:.1f}% from the average of prior extracted values"
                )
            explanation_parts.append(f"source confidence {source_confidence:.0f}%")

            candidates_by_metric.append({
                "metric_id": metric["id"],
                "metric_name": metric_name,
                "unit": unit,
                "metric_kind": metric_kind,
                "recommended_value": round(value, 4),
                "confidence": round(score, 1),
                "extraction_confidence": round(extraction_confidence, 1),
                "source_confidence": round(source_confidence, 1),
                "context": str(candidate.get("context") or ""),
                "reason_codes": list(dict.fromkeys(reason_codes)),
                "explanation": "; ".join(explanation_parts),
                "consistency_average": round(consistency_average, 4) if consistency_average is not None else None,
                "consistency_deviation_percent": round(consistency_spread, 2) if consistency_spread is not None else None,
            })

    candidates_by_metric.sort(
        key=lambda item: (
            -float(item["confidence"]),
            -float(item["extraction_confidence"]),
            str(item["metric_name"]).lower(),
        )
    )

    if not candidates_by_metric:
        return {
            "evidence_id": str(evidence_row["id"]),
            "recommendation": None,
            "alternatives": [],
            "message": (
                "No evidence value could be matched confidently to a project metric. "
                "Attach the evidence to a metric or review the file manually."
            ),
        }

    recommendation = candidates_by_metric[0]
    alternatives = [
        item for item in candidates_by_metric[1:]
        if item["metric_id"] != recommendation["metric_id"]
    ][:5]

    if alternatives and (
        float(recommendation["confidence"]) - float(alternatives[0]["confidence"]) < 4
    ):
        recommendation["selection_note"] = (
            "Top metric match is close to another metric. Review the alternatives before applying."
        )
    else:
        recommendation["selection_note"] = (
            "This is the strongest explainable match from the available evidence."
        )

    return {
        "evidence_id": str(evidence_row["id"]),
        "recommendation": recommendation,
        "alternatives": alternatives,
        "message": "Recommendation generated. Nothing was changed.",
    }


@router.post("/{project_id}/evidence/{evidence_id}/recommend")
def recommend_evidence_value(
    project_id: str,
    evidence_id: str,
    request: EvidenceRecommendationRequest,
):
    conn = get_db()
    try:
        project_exists(conn, project_id)
        row = conn.execute(
            "SELECT * FROM project_evidence WHERE project_id = ? AND CAST(id AS TEXT) = ? LIMIT 1",
            (project_id, str(evidence_id)),
        ).fetchone()
        if row is None:
            raise HTTPException(status_code=404, detail="Evidence record not found")

        return _recommend_evidence_metrics(
            conn,
            project_id,
            row,
            request.metric_id,
        )
    finally:
        conn.close()


@router.post("/{project_id}/evidence/{evidence_id}/extract")
def extract_evidence_value(project_id: str, evidence_id: str, request: EvidenceExtractRequest):
    conn = get_db()
    try:
        project_exists(conn, project_id)
        row = conn.execute(
            "SELECT * FROM project_evidence WHERE project_id = ? AND CAST(id AS TEXT) = ? LIMIT 1",
            (project_id, str(evidence_id)),
        ).fetchone()
        if row is None:
            raise HTTPException(status_code=404, detail="Evidence record not found")

        metric_id = request.metric_id or (str(row["metric_id"]) if row["metric_id"] is not None else "")
        metric = metric_belongs_to_project(conn, project_id, metric_id) if metric_id else None
        metric_name = str(metric["name"] or "") if metric else ""
        unit = str(metric["unit"] or "") if metric else ""
        target = float(metric["target"] or 0) if metric else None
        baseline = float(metric["baseline"] or 0) if metric else None
        metric_kind = str(metric["kind"] or "") if metric else ""
        path = evidence_storage_path(project_id, row)
        candidates = extract_evidence_candidates(path, metric_name, unit, target, baseline, metric_kind)
        if not candidates:
            return {"evidence_id": evidence_id, "candidates": [], "message": "No numeric candidates were found."}

        selected = candidates[0]
        if request.selected_value is not None:
            selected = min(candidates, key=lambda candidate: abs(float(candidate["value"]) - float(request.selected_value)))
        if request.apply:
            if metric is None:
                raise HTTPException(status_code=400, detail="Select a metric before applying an extracted value.")
            now = now_iso()
            observed_at = normalize_observed_at(row["observed_at"] if "observed_at" in row.keys() else None) or infer_observed_at_from_name(path.name)
            conn.execute("""
                UPDATE project_evidence
                SET extracted_value = ?, extracted_unit = ?, confidence = ?, metric_id = ?,
                    observed_at = COALESCE(observed_at, ?), applied_at = ?,
                    notes = CASE WHEN notes = '' THEN ? ELSE notes END
                WHERE CAST(id AS TEXT) = ? AND project_id = ?
            """, (selected["value"], unit, selected["confidence"], metric_id, observed_at, now,
                  f"Automatically extracted from {path.name}; reviewed and applied.", str(row["id"]), project_id))
            current = reconcile_metric_from_evidence(conn, project_id, metric_id)
            sync_project_progress(conn, project_id)
            conn.commit()
            applied_as_current = current is not None and str(current["evidence_id"]) == str(row["id"])
            return {
                "evidence_id": evidence_id, "applied": True, "selected": selected,
                "applied_as_current": applied_as_current,
                "message": ("Applied as the current metric observation." if applied_as_current else
                            "Saved as historical evidence; a newer reviewed observation remains current."),
                "project": serialize_project(conn, project_exists(conn, project_id))
            }

        return {"evidence_id": evidence_id, "applied": False, "selected": selected, "candidates": candidates}
    finally:
        conn.close()


@router.get("/{project_id}/evidence/file/{file_token}")
def download_project_evidence_file(project_id: str, file_token: str):
    from fastapi.responses import FileResponse

    conn = get_db()
    try:
        project_exists(conn, project_id)
        reference = evidence_file_reference(project_id, file_token)
        row = conn.execute(
            """
            SELECT * FROM project_evidence
            WHERE project_id = ?
              AND reference = ?
            ORDER BY id DESC
            LIMIT 1
            """,
            (project_id, reference),
        ).fetchone()

        if row is None:
            raise HTTPException(status_code=404, detail="Evidence file not found")

        filename = safe_evidence_filename(str(row["source_name"] or "evidence"))
        path = UPLOAD_DIR / project_id / f"{file_token}_{filename}"

        if not path.exists() or not path.is_file():
            raise HTTPException(status_code=404, detail="Evidence file is missing from storage")

        return FileResponse(path, filename=filename)
    finally:
        conn.close()


@router.post("/{project_id}/evidence/upload", status_code=201)
async def upload_project_evidence(
    project_id: str,
    file: UploadFile = File(...),
    metric_id: str = Form(""),
    evidence_type: EvidenceType = Form("document"),
    observed_at: str = Form(""),
    confidence: float = Form(80),
    extracted_value: Optional[float] = Form(None),
    notes: str = Form(""),
):
    conn = get_db()
    stored_path = None

    try:
        project_exists(conn, project_id)

        metric_value = metric_id.strip() or None
        if metric_value is not None:
            metric = metric_belongs_to_project(conn, project_id, metric_value)
            if metric is None:
                raise HTTPException(status_code=404, detail="Metric not found for this project")

        if not file.filename:
            raise HTTPException(status_code=400, detail="No file selected")
        if confidence < 0 or confidence > 100:
            raise HTTPException(status_code=400, detail="Confidence must be between 0 and 100")

        safe_name = safe_evidence_filename(file.filename)
        token = uuid4().hex
        project_dir = UPLOAD_DIR / project_id
        project_dir.mkdir(parents=True, exist_ok=True)
        stored_path = project_dir / f"{token}_{safe_name}"

        with stored_path.open("wb") as output:
            while True:
                chunk = await file.read(1024 * 1024)
                if not chunk:
                    break
                output.write(chunk)

        created_at = now_iso()
        reference = evidence_file_reference(project_id, token)

        evidence_values = {
            "id": token,
            "project_id": project_id,
            "metric_id": metric_value,
            "evidence_type": evidence_type,
            "source_name": file.filename,
            "title": file.filename,
            "reference": reference,
            "source_reference": reference,
            "extracted_value": None,
            "extracted_unit": "",
            "confidence": confidence,
            "notes": notes.strip(),
            "observed_at": normalize_observed_at(observed_at) or infer_observed_at_from_name(file.filename),
            "applied_at": created_at if extracted_value is not None else None,
            "is_current": 0,
            "created_at": created_at,
        }

        evidence_id = insert_row_compatible(conn, "project_evidence", evidence_values)

        if metric_value is not None and extracted_value is not None:
            reconcile_metric_from_evidence(conn, project_id, metric_value)

        sync_project_progress(conn, project_id)
        conn.commit()

        return {
            "message": "Evidence file uploaded successfully",
            "evidence_id": evidence_id,
            "filename": file.filename,
            "reference": reference,
            "project": serialize_project(conn, project_exists(conn, project_id)),
        }
    except HTTPException:
        conn.rollback()
        if stored_path and stored_path.exists():
            stored_path.unlink(missing_ok=True)
        raise
    except Exception:
        conn.rollback()
        if stored_path and stored_path.exists():
            stored_path.unlink(missing_ok=True)
        raise
    finally:
        await file.close()
        conn.close()


@router.get("/{project_id}/evidence")
def get_project_evidence(project_id: str):
    conn = get_db()
    try:
        project_exists(conn, project_id)
        rows = conn.execute("""
            SELECT e.*, m.name AS metric_name
            FROM project_evidence e
            LEFT JOIN project_metrics m ON m.id = e.metric_id
            WHERE e.project_id = ?
            ORDER BY e.id DESC
        """, (project_id,)).fetchall()
        return {"project_id": project_id, "evidence": [dict(row) for row in rows]}
    finally:
        conn.close()


@router.post("/{project_id}/evidence", status_code=201)
def add_project_evidence(project_id: str, evidence: EvidenceCreate):
    conn = get_db()
    try:
        project_exists(conn, project_id)
        metric = None
        if evidence.metric_id is not None:
            metric = metric_belongs_to_project(conn, project_id, evidence.metric_id)
            if metric is None:
                raise HTTPException(status_code=404, detail="Metric not found for this project")

        created_at = now_iso()
        evidence_id = uuid4().hex
        evidence_values = {
            "id": evidence_id,
            "project_id": project_id,
            "metric_id": evidence.metric_id,
            "evidence_type": evidence.evidence_type,
            "source_name": evidence.source_name.strip(),
            "title": evidence.source_name.strip() or "Manual evidence",  # legacy v1
            "reference": evidence.reference.strip(),
            "source_reference": evidence.reference.strip(),             # legacy v1
            "extracted_value": evidence.extracted_value,
            "extracted_unit": "",                                       # legacy v1
            "confidence": evidence.confidence,
            "notes": evidence.notes.strip(),
            "observed_at": normalize_observed_at(evidence.observed_at),
            "applied_at": created_at if evidence.extracted_value is not None else None,
            "is_current": 0,
            "created_at": created_at,
        }
        evidence_columns = table_columns(conn, "project_evidence")
        if evidence_columns.get("id"):
            id_meta = evidence_columns["id"]
            id_type = str(id_meta.get("type") or "").upper()
            if id_meta.get("pk") == 1 and "INT" in id_type:
                evidence_values.pop("id", None)

        evidence_id = insert_row_compatible(conn, "project_evidence", evidence_values)

        if evidence.metric_id is not None and evidence.extracted_value is not None:
            reconcile_metric_from_evidence(conn, project_id, evidence.metric_id)

        sync_project_progress(conn, project_id)
        conn.commit()
        return {
            "message": "Evidence recorded successfully",
            "evidence_id": evidence_id,
            "project": serialize_project(conn, project_exists(conn, project_id)),
        }
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


@router.delete("/{project_id}/evidence/{evidence_id}")
def delete_project_evidence(project_id: str, evidence_id: str):
    conn = get_db()
    try:
        project_exists(conn, project_id)
        row = conn.execute(
            "SELECT * FROM project_evidence WHERE id = ? AND project_id = ?",
            (evidence_id, project_id),
        ).fetchone()
        if row is None:
            raise HTTPException(status_code=404, detail="Evidence record not found")

        reference = str(row["reference"] or "")
        source_reference = str(row["source_reference"] or "") if "source_reference" in row.keys() else ""
        file_reference = reference if "/evidence/file/" in reference else source_reference
        if "/evidence/file/" in file_reference:
            token = file_reference.rsplit("/evidence/file/", 1)[-1].split("/", 1)[0]
            filename = safe_evidence_filename(str(row["source_name"] or "evidence"))
            stored_path = UPLOAD_DIR / project_id / f"{token}_{filename}"
            if stored_path.exists():
                stored_path.unlink(missing_ok=True)

        affected_metric_id = row["metric_id"]
        conn.execute(
            "DELETE FROM project_evidence WHERE id = ? AND project_id = ?",
            (evidence_id, project_id),
        )
        if affected_metric_id is not None:
            reconcile_metric_from_evidence(conn, project_id, affected_metric_id)
        sync_project_progress(conn, project_id)
        conn.commit()
        return {
            "message": "Evidence deleted successfully",
            "project": serialize_project(conn, project_exists(conn, project_id)),
        }
    except HTTPException:
        conn.rollback()
        raise
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


@router.put("/{project_id}/metrics/{metric_id}/value")
def update_metric_value(project_id: str, metric_id: str, update: MetricValueUpdate):
    conn = get_db()
    try:
        project_exists(conn, project_id)
        metric = metric_belongs_to_project(conn, project_id, metric_id)
        if metric is None:
            raise HTTPException(status_code=404, detail="Metric not found for this project")

        timestamp = now_iso()
        metric_columns = table_columns(conn, "project_metrics")
        update_values = {
            "current_value": update.current_value,
            "expected_value": update.expected_value,
            "last_updated": timestamp,
        }
        usable = {key: value for key, value in update_values.items() if key in metric_columns}
        if usable:
            set_clause = ", ".join(f"{key} = ?" for key in usable)
            conn.execute(
                f"UPDATE project_metrics SET {set_clause} WHERE id = ?",
                [*usable.values(), metric_id],
            )

        evidence_values = {
            "id": uuid4().hex,
            "project_id": project_id,
            "metric_id": metric_id,
            "evidence_type": update.evidence_type,
            "source_name": update.source_name.strip(),
            "title": update.source_name.strip() or "Metric update",
            "reference": update.reference.strip(),
            "source_reference": update.reference.strip(),
            "extracted_value": update.current_value,
            "extracted_unit": str(metric["unit"] or ""),
            "confidence": update.confidence,
            "notes": update.notes.strip(),
            "observed_at": timestamp[:10],
            "applied_at": timestamp,
            "is_current": 1,
            "created_at": timestamp,
        }
        evidence_columns = table_columns(conn, "project_evidence")
        if evidence_columns.get("id"):
            id_meta = evidence_columns["id"]
            id_type = str(id_meta.get("type") or "").upper()
            if id_meta.get("pk") == 1 and "INT" in id_type:
                evidence_values.pop("id", None)
        insert_row_compatible(conn, "project_evidence", evidence_values)
        reconcile_metric_from_evidence(conn, project_id, metric_id)

        sync_project_progress(conn, project_id)
        conn.commit()
        row = project_exists(conn, project_id)
        return {"message": "Metric updated successfully", "project": serialize_project(conn, row)}
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


@router.get("/{project_id}")
def get_project(project_id: str):
    conn = get_db()
    try:
        row = project_exists(conn, project_id)
        if row["progress_model"] == "universal":
            sync_project_progress(conn, project_id)
            conn.commit()
            row = project_exists(conn, project_id)
        return serialize_project(conn, row)
    finally:
        conn.close()


@router.post("/", status_code=201)
def create_project(project_data: ProjectCreate):
    conn = get_db()
    try:
        highest = conn.execute("""
            SELECT MAX(CAST(SUBSTR(id,5) AS INTEGER)) FROM projects
        """).fetchone()[0] or 0
        project_id = f"PRJ-{highest + 1:03d}"

        project_type = normalize_project_type(project_data.project_type, project_data.sector)
        progress_model = project_data.progress_model

        conn.execute("""
            INSERT INTO projects(
                id,name,sector,state,location,approved_cost,current_expenditure,
                physical_progress,expected_progress,risk_score,delay_risk,cost_risk,
                status,project_type,description,start_date,target_date,progress_model
            ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
        """, (
            project_id, project_data.name.strip(), project_data.sector.strip(),
            project_data.state.strip(), project_data.location.strip(),
            project_data.approved_cost, project_data.current_expenditure,
            project_data.physical_progress, project_data.expected_progress,
            0, "Low", "Low", project_data.status,
            project_type, project_data.description.strip(),
            project_data.start_date, project_data.target_date, progress_model
        ))

        insert_objectives_and_metrics(conn, project_id, project_data.objectives)

        # If universal metrics exist, their weighted result becomes the source
        # of truth for physical progress and risk.
        if project_data.objectives:
            sync_project_progress(conn, project_id)
        else:
            risk = calculate_risk(
                project_data.approved_cost, project_data.current_expenditure,
                project_data.physical_progress, project_data.expected_progress
            )
            conn.execute("""
                UPDATE projects SET risk_score=?, delay_risk=?, cost_risk=?
                WHERE id=?
            """, (risk["risk_score"], risk["delay_risk"], risk["cost_risk"], project_id))

        conn.commit()
        row = project_exists(conn, project_id)
        return {"message": "Project created successfully", "project": serialize_project(conn, row)}
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


@router.put("/{project_id}")
def update_project(project_id: str, project_data: ProjectCreate):
    conn = get_db()
    try:
        existing = project_exists(conn, project_id)

        project_type = normalize_project_type(project_data.project_type, project_data.sector)
        conn.execute("""
            UPDATE projects SET
                name=?, sector=?, state=?, location=?,
                approved_cost=?, current_expenditure=?,
                physical_progress=?, expected_progress=?,
                status=?, project_type=?, description=?,
                start_date=?, target_date=?, progress_model=?
            WHERE id=?
        """, (
            project_data.name.strip(), project_data.sector.strip(),
            project_data.state.strip(), project_data.location.strip(),
            project_data.approved_cost, project_data.current_expenditure,
            project_data.physical_progress, project_data.expected_progress,
            project_data.status, project_type, project_data.description.strip(),
            project_data.start_date, project_data.target_date,
            project_data.progress_model, project_id
        ))

        if project_data.objectives:
            conn.execute("DELETE FROM project_objectives WHERE project_id = ?", (project_id,))
            insert_objectives_and_metrics(conn, project_id, project_data.objectives)
            sync_project_progress(conn, project_id)
        else:
            risk = calculate_risk(
                project_data.approved_cost, project_data.current_expenditure,
                project_data.physical_progress, project_data.expected_progress
            )
            conn.execute("""
                UPDATE projects SET risk_score=?, delay_risk=?, cost_risk=? WHERE id=?
            """, (risk["risk_score"], risk["delay_risk"], risk["cost_risk"], project_id))

        conn.commit()
        row = project_exists(conn, project_id)
        return {"message": "Project updated successfully", "project": serialize_project(conn, row)}
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


@router.delete("/{project_id}")
def delete_project(project_id: str):
    conn = get_db()
    try:
        row = project_exists(conn, project_id)
        deleted = serialize_project(conn, row)
        conn.execute("DELETE FROM projects WHERE id = ?", (project_id,))
        conn.commit()
        return {"message": "Project deleted successfully", "project": deleted}
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()
